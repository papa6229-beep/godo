"""바나나몰 엑셀 파서 — DESIGN.md 2.1.

68열이지만 열 **순서를 믿지 않는다.** 헤더 이름으로 찾는다. 순서로 읽으면
다운로드 양식이 한 번만 바뀌어도 조용히 엉뚱한 값을 집는다.

찾지 못한 헤더는 조용히 넘어가지 않고 이름을 들고 올라온다. 1000개를 돌리기 전에
"이 파일에는 이런 헤더가 있다"를 사람이 볼 수 있어야 한다.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field

#: 우리가 쓰는 이름 → 엑셀에 나올 법한 헤더 후보들.
FIELDS: dict[str, tuple[str, ...]] = {
    "code": ("상품번호", "상품코드", "goodsno", "품번"),
    "name": ("상품명", "제품명", "goodsnm"),
    "brand": ("브랜드", "제조사", "brand"),
    "maker": ("거래처", "공급사", "제조사명"),
    "price": ("판매가", "판매가격", "price"),
    "cost": ("원가", "정가", "소비자가"),
    "category": ("카테고리", "분류"),
    "thumb": ("대표이미지url", "대표이미지", "목록이미지url", "목록이미지"),
    "body": ("상세설명", "상세정보", "상품상세", "detail"),
    "url": ("상품url", "상품링크"),
}

OPTION_HEADERS = ("옵션1", "옵션2", "옵션3", "옵션")


def norm(s) -> str:
    """헤더 비교용 정규화 — 공백·괄호·기호를 지우고 소문자로."""
    return re.sub(r"[\s()\[\]/_\-.]", "", str(s or "")).lower()


@dataclass
class Option:
    name: str
    values: list[str] = field(default_factory=list)
    #: 원본에 붙어 있던 번호(`01`). 없으면 빈 칸. values 와 길이가 같다.
    numbers: list[str] = field(default_factory=list)


@dataclass
class Row:
    code: str = ""
    name: str = ""
    brand: str = ""
    maker: str = ""
    price: str = ""
    cost: str = ""
    category: str = ""
    thumb: str = ""
    body: str = ""
    url: str = ""
    options: list[Option] = field(default_factory=list)
    #: 엑셀에서 읽은 원본 행 전체. 못 쓴 열을 나중에 확인할 수 있게 남긴다.
    raw: dict = field(default_factory=dict)

    @property
    def option_values(self) -> list[str]:
        return [v for o in self.options for v in o.values]

    @property
    def option_numbers(self) -> list[str]:
        """option_values 와 같은 자리의 원본 번호. 없던 자리는 빈 칸."""
        return [n for o in self.options for n in o.numbers]


#: 코드 한 도막 — 영대문자·숫자·하이픈·밑줄만. 한글이나 소문자가 섞이면 이름이다.
_TOKEN = r"[A-Z0-9][A-Z0-9_\-]*"
_GROUP = rf"{_TOKEN}(?:\s*/\s*{_TOKEN})*"
#: 괄호 코드 `(EGG-013)` `(4582236080170)` `(OH-3584/4580664902224)` `(DJ)`.
#: 끝이 아니라 **어디 있든** 뗀다 — `유니 다이아몬드 (UNI-002) (화이트)` 처럼
#: 뒤에 이름 괄호가 한 번 더 붙는 경우가 있다.
_PAREN_CODE = re.compile(rf"\s*\(\s*{_GROUP}\s*\)")
#: 하이픈 꼬리 `- OH-3650/4570099420325` `- 4526374570674`
_DASH_TAIL = re.compile(rf"\s*[-–]\s*({_GROUP})\s*$")
#: `- solvemen029/4580490010322` — 코드 쪽에 소문자가 섞여도 뒤가 바코드면 코드다.
_SLASH_BARCODE = re.compile(r"\s*[-–]\s*\S+\s*/\s*\d{8,14}\s*$")


def _strip_tail(value: str) -> str:
    """물류용 꼬리를 뗀다. 이름이면 안 뗀다."""
    out = _PAREN_CODE.sub("", value).strip() or value.strip()
    for _ in range(3):
        m = _SLASH_BARCODE.search(out) or None
        if m and out[: m.start()].strip():
            out = out[: m.start()].strip()
            continue
        m = _DASH_TAIL.search(out)
        # 하이픈 뒤는 숫자가 있어야 코드로 본다. `블랙 - XL` 의 XL 을 지우지 않기 위해서다.
        if m and any(c.isdigit() for c in m.group(1)) and out[: m.start()].strip():
            out = out[: m.start()].strip()
            continue
        break
    return out


def strip_codes(values: list[str]) -> list[str]:
    """옵션값에서 모델코드·바코드를 뗀다 — **한 축을 통째로 보고 정한다.**

    엑셀 옵션값의 60%(892개 중 532개)에 `- OH-3650/4570099420325` 같은 꼬리가 붙어
    있다. 그래서 상세페이지 말머리 `[부부장 유아]` 와 글자로 대조하면 **하나도 안 맞는다** —
    실물 49개에서 옵션 상품 17개가 전부 사진 없는 빈 카드로 나왔다.

    떼고 나서 **빈 값이 되거나 서로 겹치면 아예 안 뗀다.** `블랙(L)` / `블랙(XL)` 처럼
    꼬리가 옵션을 가르는 유일한 표시일 수 있기 때문이다. 그런 자리는 겹침으로 드러나므로
    따로 판정할 것이 없다 — 892개 실측에서는 빈 값도 겹침도 0이었다.
    """
    out = [_strip_tail(v) for v in values]
    if any(not v for v in out) or len(set(out)) != len(set(values)):
        return list(values)
    return out


def parse_option(cell: str) -> Option | None:
    """`옵션명=값,값,값` 형식을 푼다 (2.1).

    7장 — 원본 접두 번호(`01. `)는 값에서 벗기되 **버리지는 않는다.**
    그대로 두면 대조가 안 되고(캡션 쪽도 벗겨서 맞춘다), 버리면 손님이 주문할 때
    몇 번 옵션인지 알 수 없다. 값과 번호를 나란히 들고 간다.
    """
    text = str(cell or "").strip()
    if not text:
        return None
    name, sep, rest = text.partition("=")
    if not sep:
        name, rest = "", text
    values, numbers = [], []
    for v in rest.split(","):
        m = re.match(r"^\s*(\d+)\s*[.)]\s*", v.strip())
        text = v.strip()[m.end():].strip() if m else v.strip()
        if text:
            values.append(text)
            numbers.append(m.group(1) if m else "")
    if not values:
        return None
    return Option(name=name.strip(), values=strip_codes(values), numbers=numbers)


def _header_map(header: list) -> dict[str, int]:
    """헤더 행 → {우리 이름: 열 번호}."""
    seen = {norm(h): i for i, h in enumerate(header) if str(h or "").strip()}
    out: dict[str, int] = {}
    for key, candidates in FIELDS.items():
        for cand in candidates:
            i = seen.get(norm(cand))
            if i is not None:
                out[key] = i
                break
    return out


def _find_header_row(rows: list[list], limit: int = 10) -> int:
    """헤더가 첫 줄이 아닐 수 있다. 아는 이름이 가장 많이 걸리는 줄을 고른다."""
    best, best_hits = 0, -1
    for i, row in enumerate(rows[:limit]):
        hits = len(_header_map(row))
        if hits > best_hits:
            best, best_hits = i, hits
    return best


@dataclass
class Sheet:
    rows: list[Row]
    headers: list[str]
    #: 찾지 못한 항목. UI 에 그대로 보여준다.
    missing: list[str]


def load(path_or_bytes) -> Sheet:
    """엑셀(.xlsx)을 읽어 행 목록으로."""
    import io

    import openpyxl

    src = io.BytesIO(path_or_bytes) if isinstance(path_or_bytes, (bytes, bytearray)) else path_or_bytes
    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    grid = [list(r) for r in ws.iter_rows(values_only=True)]
    if not grid:
        raise ValueError("빈 시트다")

    hi = _find_header_row(grid)
    header = [str(h or "").strip() for h in grid[hi]]
    cols = _header_map(grid[hi])
    missing = [k for k in ("code", "name", "body") if k not in cols]

    opt_cols = [
        i for i, h in enumerate(grid[hi])
        if any(norm(h) == norm(o) for o in OPTION_HEADERS)
    ]

    rows: list[Row] = []
    for raw in grid[hi + 1 :]:
        if not any(str(c or "").strip() for c in raw):
            continue
        get = lambda k: str(raw[cols[k]] or "").strip() if k in cols and cols[k] < len(raw) else ""  # noqa: E731
        row = Row(
            code=get("code"), name=get("name"), brand=get("brand"), maker=get("maker"),
            price=get("price"), cost=get("cost"), category=get("category"),
            thumb=get("thumb"), body=get("body"), url=get("url"),
            raw={header[i]: raw[i] for i in range(min(len(header), len(raw))) if header[i]},
        )
        for i in opt_cols:
            if i < len(raw):
                opt = parse_option(raw[i])
                if opt:
                    row.options.append(opt)
        if row.code or row.name:
            rows.append(row)

    wb.close()
    return Sheet(rows=rows, headers=[h for h in header if h], missing=missing)
